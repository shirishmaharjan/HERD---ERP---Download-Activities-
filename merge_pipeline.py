#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
merge_pipeline.py
=================

Headless version of the ERP download + merge workflow, used by the Streamlit
app (``app.py``).  It performs the same steps as the command-line script but
returns data instead of printing:

    login -> read ERP projects -> match Sheet2 rows -> download exports
          -> merge -> write Sheet3 (+ Project Map) into the workbook

The workbook passed in is modified in place, so callers should hand it a copy
(``app.py`` writes the uploaded file to a temporary folder first).
"""

from __future__ import annotations

import urllib.parse
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

import openpyxl
import pandas as pd

try:  # module name in this workspace
    from erp_export_merge import (
        BASE_URL,
        ErpClient,
        ErpError,
        _first,
        download_export,
        match_projects,
        read_project_names,
        read_table,
        save_workbook,
        scrape_projects,
        write_sheet,
    )
except ImportError:  # deployed next to the CLI copy named erp_export.py
    from erp_export import (  # type: ignore[no-redef]
        BASE_URL,
        ErpClient,
        ErpError,
        _first,
        download_export,
        match_projects,
        read_project_names,
        read_table,
        save_workbook,
        scrape_projects,
        write_sheet,
    )

Progress = Callable[[int, int, str], None]
Logger = Callable[[str], None]


@dataclass
class RunResult:
    output_path: Path
    total_projects: int
    ok_projects: int
    merged_rows: int
    mapping: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def run_merge(
    workbook_path: str | Path,
    username: str,
    password: str,
    sheet: str = "Sheet2",
    out_sheet: str = "Sheet3",
    base_url: str = BASE_URL,
    exports_dir: str | Path | None = None,
    progress: Progress | None = None,
    log: Logger | None = None,
    insecure: bool = False,
) -> RunResult:
    workbook = Path(workbook_path)
    exports = Path(exports_dir) if exports_dir else workbook.parent / "erp_exports"
    exports.mkdir(parents=True, exist_ok=True)

    def step(done: int, total: int, message: str) -> None:
        if progress:
            progress(done, total, message)

    def note(message: str) -> None:
        if log:
            log(message)

    names, _header_row, _column = read_project_names(workbook, sheet)
    if not names:
        raise ErpError(f"No project names found in sheet {sheet!r}.")

    total_steps = len(names) + 3
    step(0, total_steps, "Signing in to the ERP ...")
    client = ErpClient(base_url, insecure=insecure)
    client.login(username, password)

    step(1, total_steps, "Reading the ERP project list ...")
    erp_projects = scrape_projects(client)
    if not erp_projects:
        raise ErpError("Could not read any project from the ERP project list.")
    note(f"ERP returned {len(erp_projects)} projects; Sheet2 has {len(names)}.")

    matches = match_projects(names, erp_projects)
    token = _first(
        client.get("/projects").text,
        r'name="csrf-token"\s+content="([^"]+)"',
        r'name="_token"\s+value="([^"]+)"',
    )
    referer = urllib.parse.urljoin(client.base + "/", "/projects")

    downloaded: list[tuple] = []
    status_rows: list[dict] = []
    warnings: list[str] = []

    for index, match in enumerate(matches, start=1):
        row = {
            "Project (as in Excel)": match.full_name,
            "Short": match.short,
            "ERP Project": match.project.name if match.project else "",
            "ERP ID": match.project.pid if match.project else "",
            "Status": "unmatched",
            "File": "",
            "Rows": "",
            "Note": match.note,
        }
        if match.project is None:
            status_rows.append(row)
            warnings.append(f"{match.short}: not found in the ERP ({match.note})")
            continue

        step(index + 1, total_steps, f"Downloading {match.short} ...")
        try:
            path, message = download_export(
                client, match.project, match.short, exports,
                referer=referer, token=token, reuse=False, keep_html=True,
            )
        except Exception as exc:  # noqa: BLE001 - one project must not stop the rest
            row.update(Status="failed", Note=f"{type(exc).__name__}: {exc}")
            status_rows.append(row)
            warnings.append(f"{match.short}: download failed ({exc})")
            continue

        if path is None:
            row.update(Status="failed", Note=message)
            status_rows.append(row)
            warnings.append(f"{match.short}: {message}")
            continue

        row.update(Status="ok", File=path.name)
        downloaded.append((match, path))
        status_rows.append(row)

    step(len(names) + 2, total_steps, "Merging files ...")
    frames = []
    for match, path in downloaded:
        try:
            frame = read_table(path)
        except Exception as exc:  # noqa: BLE001
            warnings.append(f"{match.short}: could not read {path.name} ({exc})")
            continue
        if frame.empty:
            warnings.append(f"{match.short}: export contained no rows")
            continue
        frame.insert(0, "Project (Full)", match.project.name)
        frame.insert(0, "Project (Short)", match.short)
        frames.append(frame)
        for row in status_rows:
            if row["Short"] == match.short and row["Status"] == "ok":
                row["Rows"] = len(frame)

    if not frames:
        raise ErpError("No usable rows were read from the downloaded exports.")

    merged = pd.concat(frames, ignore_index=True, sort=False).dropna(how="all")

    workbook_obj = openpyxl.load_workbook(workbook, data_only=False, keep_links=True)
    if out_sheet in workbook_obj.sheetnames:
        del workbook_obj[out_sheet]
    write_sheet(workbook_obj.create_sheet(out_sheet), merged)

    map_name = "Project Map"
    if map_name in workbook_obj.sheetnames:
        del workbook_obj[map_name]
    map_frame = pd.DataFrame(status_rows, columns=[
        "Project (as in Excel)", "Short", "ERP Project", "ERP ID",
        "Status", "File", "Rows", "Note",
    ])
    write_sheet(workbook_obj.create_sheet(map_name), map_frame)

    step(total_steps, total_steps, "Saving workbook ...")
    save_workbook(workbook_obj, workbook)

    ok = sum(1 for row in status_rows if row["Status"] == "ok")
    return RunResult(
        output_path=workbook,
        total_projects=len(matches),
        ok_projects=ok,
        merged_rows=len(merged),
        mapping=status_rows,
        warnings=warnings,
    )
