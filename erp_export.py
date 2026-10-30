#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
erp_export_merge.py
===================

Download the "activities" export of every project listed in an Excel sheet from
the HERD ERP (https://erp.herdint.com) and merge all of them into a single new
sheet (default: ``Sheet3``) inside the same workbook.

What it does
------------
1. Reads the project list from the workbook (default sheet ``Sheet2``).
2. Extracts the short name from the last bracket of every project, e.g.
   ``CHORUS: Adolescent Mental Health (CHORUS IF - AMH)`` -> ``CHORUS IF - AMH``.
3. Logs in to the ERP (username / password from the environment or a prompt).
4. Scrapes ``/projects`` (all pages) to map every project name to its numeric id.
5. Matches each Excel project to an ERP project (short code first, then fuzzy
   name similarity) and prints the mapping so it can be reviewed.
6. Opens ``/projects/<id>/dashboard`` -> "Activities" -> finds the "Export"
   control (link or form, GET or POST) and downloads the file.
7. Merges every downloaded file into one table, adds ``Project (Short)`` and
   ``Project (Full)`` columns, and writes it to ``Sheet3`` of the same workbook
   (a timestamped backup of the workbook is created first).

Usage
-----
    set HERD_ERP_USER=your.user
    set HERD_ERP_PASS=your-password
    python erp_export_merge.py --workbook "C:\\path\\mel raw data.xlsx"

    # only show the project matching, download nothing
    python erp_export_merge.py --list-only

    # use files that were already downloaded
    python erp_export_merge.py --reuse

Requirements
------------
    pip install pandas openpyxl

Only the standard library is used for the HTTP part, so ``requests`` is not
needed.  If the ERP uses a self-signed certificate, add ``--insecure``.
"""

from __future__ import annotations

import argparse
import io
import json
import math
import os
import re
import shutil
import ssl
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from difflib import SequenceMatcher
from html import unescape
from http.cookiejar import CookieJar
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- #
# Defaults
# --------------------------------------------------------------------------- #

BASE_URL = "https://erp.herdint.com"
VERSION = "2026-10-30.7 (save retry + verification)"
DEFAULT_WORKBOOK = (
    r"C:\Users\Shirish Maharjan\OneDrive - HERD\Herd\M&E\Data\29.10.2026\mel raw data.xlsx"
)
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36"
)
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(vertical="center", horizontal="left")


def log(message: str = "") -> None:
    print(message, flush=True)


# --------------------------------------------------------------------------- #
# Small text helpers
# --------------------------------------------------------------------------- #

def clean_text(value: str) -> str:
    value = unescape(value or "")
    value = value.replace("\xa0", " ")
    return re.sub(r"\s+", " ", value).strip()


def strip_tags(value: str) -> str:
    return clean_text(re.sub(r"<[^>]+>", " ", value or ""))


def short_name(full_name: str) -> str:
    """Return the text inside the last bracket of a project name.

    ``CHORUS: Adolescent Mental Health (CHORUS IF - AMH)`` -> ``CHORUS IF - AMH``
    ``CHORUS I: Strengthening Urban Health Systems (CHORUS I`` -> ``CHORUS I``
    """
    text = clean_text(full_name)
    match = re.search(r"\(([^()]*)\)\s*$", text)
    if match:
        return match.group(1).strip()
    # Missing closing bracket, e.g. "(CHORUS I"
    if text.count("(") == 1 and ")" not in text[text.index("(") :]:
        return text[text.index("(") + 1 :].strip()
    return text


def norm(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (value or "").lower()).strip()


def safe_filename(value: str, limit: int = 60) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", value or "").strip("._")
    return (value or "project")[:limit]


# --------------------------------------------------------------------------- #
# HTTP client (standard library only)
# --------------------------------------------------------------------------- #

class ErpError(RuntimeError):
    pass


class AuthError(ErpError):
    pass


@dataclass
class Response:
    status: int
    url: str
    headers: dict
    content: bytes

    @property
    def text(self) -> str:
        charset = "utf-8"
        match = re.search(r"charset=([\w-]+)", self.headers.get("content-type", ""))
        if match:
            charset = match.group(1)
        try:
            return self.content.decode(charset, "replace")
        except LookupError:
            return self.content.decode("utf-8", "replace")

    def header(self, name: str, default: str = "") -> str:
        return self.headers.get(name.lower(), default)


class ErpClient:
    def __init__(self, base_url: str, timeout: int = 120, retries: int = 3,
                 insecure: bool = False, delay: float = 0.0):
        self.base = base_url.rstrip("/")
        self.timeout = timeout
        self.retries = retries
        self.delay = delay
        self.jar = CookieJar()
        context = ssl.create_default_context()
        if insecure:
            context.check_hostname = False
            context.verify_mode = ssl.CERT_NONE
        handlers = [
            urllib.request.HTTPCookieProcessor(self.jar),
            urllib.request.HTTPSHandler(context=context),
        ]
        self.opener = urllib.request.build_opener(*handlers)
        self.opener.addheaders = [
            ("User-Agent", USER_AGENT),
            ("Accept", "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8"),
            ("Accept-Language", "en-US,en;q=0.9"),
        ]
        self.username: str | None = None

    # -- low level ---------------------------------------------------------
    def request(self, method: str, url: str, data: dict | None = None,
                headers: dict | None = None, referer: str | None = None,
                timeout: int | None = None) -> Response:
        full_url = url if url.startswith("http") else urllib.parse.urljoin(self.base + "/", url)
        body = None
        request_headers = {}
        if headers:
            request_headers.update(headers)
        if referer:
            request_headers["Referer"] = referer
        if data is not None:
            body = urllib.parse.urlencode(data, doseq=True).encode("utf-8")
            request_headers["Content-Type"] = "application/x-www-form-urlencoded"

        last_error: Exception | None = None
        for attempt in range(1, self.retries + 1):
            try:
                if self.delay:
                    time.sleep(self.delay)
                request = urllib.request.Request(
                    full_url, data=body, headers=request_headers, method=method.upper()
                )
                with self.opener.open(request, timeout=timeout or self.timeout) as raw:
                    content = raw.read()
                    return Response(
                        status=raw.status,
                        url=raw.geturl(),
                        headers={k.lower(): v for k, v in raw.headers.items()},
                        content=content,
                    )
            except urllib.error.HTTPError as exc:
                content = b""
                try:
                    content = exc.read()
                except Exception:
                    pass
                if exc.code in (429, 500, 502, 503, 504) and attempt < self.retries:
                    time.sleep(1.5 * attempt)
                    last_error = exc
                    continue
                return Response(
                    status=exc.code,
                    url=full_url,
                    headers={k.lower(): v for k, v in (exc.headers or {}).items()},
                    content=content,
                )
            except (urllib.error.URLError, TimeoutError, ConnectionError, ssl.SSLError) as exc:
                last_error = exc
                if attempt < self.retries:
                    time.sleep(1.5 * attempt)
                    continue
        raise ErpError(f"{method} {full_url} failed: {last_error}")

    def get(self, url: str, **kwargs) -> Response:
        return self.request("GET", url, **kwargs)

    def post(self, url: str, data: dict, **kwargs) -> Response:
        return self.request("POST", url, data=data, **kwargs)

    # -- auth --------------------------------------------------------------
    def login(self, username: str, password: str) -> None:
        page = self.get("/login")
        if page.status >= 400:
            raise AuthError(f"Could not open the login page (HTTP {page.status}).")

        token = _first(
            page.text,
            r'name="_token"[^>]*value="([^"]+)"',
            r'value="([^"]+)"[^>]*name="_token"',
            r'name="csrf-token"\s+content="([^"]+)"',
        )
        action = _first(page.text, r'<form[^>]*action="([^"]*login[^"]*)"') or "/login"

        payload = {
            "username": username,
            "password": password,
            "remember": "on",
            "previous": self.base,
        }
        if token:
            payload["_token"] = token

        self.post(action, payload, referer=page.url)

        check = self.get("/projects")
        if _is_login_page(check):
            raise AuthError(
                "Login failed - the ERP returned the login page again. "
                "Check your username/password."
            )
        self.username = username
        log(f"Logged in as {username}.")


def _first(text: str, *patterns: str) -> str:
    for pattern in patterns:
        match = re.search(pattern, text, re.I)
        if match:
            return match.group(1)
    return ""


def _is_login_page(response: Response) -> bool:
    if response.url.rstrip("/").endswith("/login"):
        return True
    head = response.text[:6000]
    return bool(re.search(r"<title>\s*(User )?Login", head, re.I)) or (
        'name="password"' in head and 'name="username"' in head
    )


# --------------------------------------------------------------------------- #
# Project discovery on /projects
# --------------------------------------------------------------------------- #

@dataclass
class ErpProject:
    pid: str
    name: str
    short: str = ""
    source: str = ""


_ROW_RE = re.compile(r"<tr\b[^>]*>(.*?)</tr>", re.S | re.I)
_HREF_PROJECT_RE = re.compile(
    r'href="([^"]*?/projects/(\d+)(?:/(?:dashboard|activities|overview))?[^"]*)"', re.I
)


def parse_project_rows(html: str, page_url: str) -> list[ErpProject]:
    """Extract (id, name) pairs from the projects table."""
    projects: list[ErpProject] = []
    seen: dict[str, ErpProject] = {}

    for row in _ROW_RE.finditer(html):
        block = row.group(1)
        links = _HREF_PROJECT_RE.findall(block)
        if not links:
            continue
        pid = links[0][1]
        text = strip_tags(block)
        # Remove the action labels that live inside the same row.
        text = re.sub(
            r"\b(view project|view details|open project|dashboard|activities|edit|delete)\b",
            " ",
            text,
            flags=re.I,
        )
        text = clean_text(text)
        name = _best_name(block, text)
        candidate = ErpProject(pid=pid, name=name, source=page_url)
        if pid in seen:
            if len(candidate.name) > len(seen[pid].name):
                seen[pid] = candidate
        else:
            seen[pid] = candidate

    if not seen:  # fallback: no table markup
        for href, pid in _HREF_PROJECT_RE.findall(html):
            if pid not in seen:
                seen[pid] = ErpProject(pid=pid, name=f"Project {pid}", source=page_url)

    projects = list(seen.values())
    return projects


_ACTION_CELL_RE = re.compile(
    r"^(view|view project|view details|open|open project|edit|delete|details|"
    r"dashboard|activities|settings|active|inactive|completed|ongoing)$",
    re.I,
)
_NUMERIC_CELL_RE = re.compile(r"^[\d\s.,%/:-]+$")


def _best_name(block: str, row_text: str) -> str:
    """Pick the project name out of a table row.

    Taking the whole row text would glue neighbouring cells (donor, status,
    action buttons) onto the name, which breaks short-code extraction, so the
    individual ``<td>`` cells are examined first.
    """
    options: list[str] = []
    for cell in re.findall(r"<td\b[^>]*>(.*?)</td>", block, re.S | re.I):
        text = strip_tags(cell)
        if len(text) < 3 or _ACTION_CELL_RE.match(text) or _NUMERIC_CELL_RE.match(text):
            continue
        options.append(text)

    for attr in ("title", "data-original-title", "aria-label", "data-search"):
        for value in re.findall(rf'{attr}="([^"]+)"', block, re.I):
            value = clean_text(value)
            if len(value) > 3 and not re.search(r"view project|^(edit|delete)$", value, re.I):
                options.append(value)

    if options:
        return max(options, key=len)
    return row_text


def _next_page_url(html: str, page_url: str, visited: set[str]) -> str | None:
    for match in re.finditer(r'<a[^>]+href="([^"]+)"[^>]*>(.*?)</a>', html, re.S | re.I):
        href, label = match.group(1), strip_tags(match.group(2)).lower()
        if not re.search(r"(next|›|»|older|→)", label + " " + href, re.I):
            continue
        url = urllib.parse.urljoin(page_url, unescape(href))
        if url not in visited and "page=" in url:
            return url
    return None


def find_datatable_endpoint(html: str) -> str:
    """Return the ``ajax:`` URL used by the projects DataTable, if present."""
    match = re.search(r"""ajax\s*:\s*["']([^"']+)["']""", html)
    return unescape(match.group(1)) if match else ""


def _with_params(url: str, **extra: str) -> str:
    parts = urllib.parse.urlsplit(url)
    query = dict(urllib.parse.parse_qsl(parts.query, keep_blank_values=True))
    query.update(extra)
    return urllib.parse.urlunsplit(
        (parts.scheme, parts.netloc, parts.path, urllib.parse.urlencode(query), "")
    )


def _same_host(client: "ErpClient", url: str) -> str:
    """Keep the logged-in host even if the page hard-codes another one."""
    parts = urllib.parse.urlsplit(url)
    if not parts.netloc:
        return url
    base = urllib.parse.urlsplit(client.base)
    if parts.netloc.lower() != base.netloc.lower():
        return urllib.parse.urlunsplit((base.scheme, base.netloc, parts.path, parts.query, ""))
    return url


def _project_from_json_row(row: dict) -> ErpProject | None:
    if not isinstance(row, dict):
        return None
    title = clean_text(str(row.get("title") or row.get("name") or row.get("project_name") or ""))
    short = clean_text(str(row.get("short_name") or row.get("short") or ""))
    action = str(row.get("action") or "")
    pid = str(row.get("id") or row.get("project_id") or "")
    if not pid:
        match = re.search(r"/projects/(\d+)", action)
        pid = match.group(1) if match else ""
    if not pid:
        match = re.search(r"project[_-]?id=(\d+)", action, re.I)
        pid = match.group(1) if match else ""
    if not pid or not title:
        return None
    return ErpProject(pid=pid, name=title, short=short)


JSON_HEADERS = {
    "X-Requested-With": "XMLHttpRequest",
    "Accept": "application/json, text/javascript, */*; q=0.01",
}


def _datatable_columns(html: str) -> list[str]:
    """Column ``data`` keys from the page's DataTable configuration."""
    match = re.search(r"\.DataTable\s*\(\s*\{(.*?)\}\s*\)\s*;", html, re.S)
    block = match.group(1) if match else html
    return re.findall(r"""data\s*:\s*['"]([^'"]+)['"]""", block)


def _datatable_params(draw: int, start: int, length: int, columns: list[str]) -> dict:
    params = {
        "draw": str(draw),
        "start": str(start),
        "length": str(length),
        "search[value]": "",
        "search[regex]": "false",
        "order[0][column]": "0",
        "order[0][dir]": "asc",
    }
    for index, column in enumerate(columns):
        params[f"columns[{index}][data]"] = column
        params[f"columns[{index}][name]"] = column
        params[f"columns[{index}][searchable]"] = "true"
        params[f"columns[{index}][orderable]"] = "false" if column == "DT_RowIndex" else "true"
        params[f"columns[{index}][search][value]"] = ""
        params[f"columns[{index}][search][regex]"] = "false"
    return params


def _json_payload(response: Response) -> dict | None:
    text = response.text.strip()
    if not text.startswith("{"):
        return None
    try:
        payload = json.loads(text)
    except json.JSONDecodeError:
        return None
    return payload if isinstance(payload, dict) else None


def _fetch_projects_page(client: ErpClient, url: str, params: dict, start: int,
                         draw: int) -> tuple[list[ErpProject], int | None, str]:
    query = {**params, "draw": str(draw), "start": str(start)}
    response = client.get(_with_params(url, **query), headers=JSON_HEADERS)
    if _is_login_page(response):
        raise AuthError("Session expired while reading /projects - please log in again.")

    payload = _json_payload(response)
    if payload is None:
        snippet = re.sub(r"\s+", " ", response.text[:140])
        return [], None, f"HTTP {response.status} {response.header('content-type')} body={snippet!r}"

    data = payload.get("data") or payload.get("aaData") or payload.get("rows") or []
    if not isinstance(data, list):
        data = []
    total = payload.get("recordsTotal")
    try:
        total = int(total) if total is not None else None
    except (TypeError, ValueError):
        total = None
    note = f"HTTP {response.status} json rows={len(data)} total={total}"
    return list(filter(None, (_project_from_json_row(row) for row in data))), total, note


def collect_projects_from_endpoint(
    client: ErpClient, endpoint: str, columns: list[str]
) -> tuple[list[ErpProject], list[str]]:
    """Try the DataTable endpoint with progressively simpler parameter sets."""
    diagnostics: list[str] = []

    for active in ("1", "0"):  # active projects, then inactive ones
        url = _with_params(endpoint, active=active)
        strategies: list[tuple[str, dict]] = [("bare", {})]
        if columns:
            strategies.insert(0, ("full", _datatable_params(1, 0, 200, columns)))
        strategies.append(("basic", {"draw": "1", "start": "0", "length": "200",
                                     "search[value]": ""}))

        for name, params in strategies:
            collected: list[ErpProject] = []
            start = 0
            total = None
            note = ""
            for page in range(40):
                rows, page_total, note = _fetch_projects_page(
                    client, url, params, start, page + 1
                )
                if not rows:
                    break
                collected.extend(rows)
                start += len(rows)
                total = page_total or total
                if total is not None and start >= total:
                    break
            diagnostics.append(f"active={active} {name}: {note} collected={len(collected)}")
            if collected:
                return collected, diagnostics

    return [], diagnostics


def scrape_projects(client: ErpClient, max_pages: int = 40) -> list[ErpProject]:
    projects: dict[str, ErpProject] = {}

    page = client.get("/projects")
    if _is_login_page(page):
        raise AuthError("Session expired while reading /projects - please log in again.")
    if page.status >= 400:
        raise ErpError(f"/projects returned HTTP {page.status}")

    # The listing is a server-side DataTable; its JSON carries title + short_name,
    # which is far more reliable than the rendered (empty) table body.
    endpoint = find_datatable_endpoint(page.text)
    diagnostics: list[str] = []
    columns: list[str] = []
    if endpoint:
        endpoint = _same_host(client, endpoint)
        columns = _datatable_columns(page.text)
        found, diagnostics = collect_projects_from_endpoint(client, endpoint, columns)
        for project in found:
            projects.setdefault(project.pid, project)

    if not projects:  # fallback: plain HTML table
        url: str | None = page.url
        visited: set[str] = set()
        while url and len(visited) < max_pages:
            visited.add(url)
            response = client.get(url)
            found = parse_project_rows(response.text, response.url)
            for project in found:
                if project.pid not in projects or len(project.name) > len(projects[project.pid].name):
                    projects[project.pid] = project
            url = _next_page_url(response.text, response.url, visited)

    if not projects:
        _dump_projects_debug(page.text, endpoint, diagnostics, columns)

    return list(projects.values())


def _dump_projects_debug(html: str, endpoint: str, diagnostics: list[str],
                         columns: list[str]) -> None:
    try:
        debug_dir = Path("erp_debug")
        debug_dir.mkdir(exist_ok=True)
        (debug_dir / "projects_page.html").write_text(html, encoding="utf-8")
        lines = [
            f"endpoint: {endpoint}",
            f"columns: {columns}",
            "",
            *diagnostics,
        ]
        (debug_dir / "projects_endpoint.txt").write_text("\n".join(lines), encoding="utf-8")
        log(f"  diagnostics written to {debug_dir.resolve()}")
    except OSError as exc:
        log(f"  could not write diagnostics: {exc}")
    for line in diagnostics:
        log(f"  {line}")


# --------------------------------------------------------------------------- #
# Matching Excel rows to ERP projects
# --------------------------------------------------------------------------- #

def _contains(haystack: str, needle: str) -> bool:
    """Word-boundary containment on already-normalised strings."""
    if not haystack or not needle:
        return False
    return bool(re.search(rf"(?<![a-z0-9]){re.escape(needle)}(?![a-z0-9])", haystack))


def _stem(token: str) -> str:
    """Crude singular/plural folding so 'Systems' matches 'System'."""
    return token[:-1] if len(token) > 3 and token.endswith("s") else token


def name_similarity(short: str, full: str, erp_name: str, erp_short_field: str = "") -> float:
    s, f, e = norm(short), norm(full), norm(erp_name)
    erp_short = norm(erp_short_field) if erp_short_field else norm(short_name(erp_name))
    if s and erp_short and s == erp_short:
        return 1.0

    scores = [0.0]
    if s and erp_short:
        if s == erp_short:
            scores.append(1.0)
        elif re.match(rf"^{re.escape(s)}(?![a-z0-9])", erp_short):
            scores.append(0.96)
        elif _contains(erp_short, s) or _contains(s, erp_short):
            scores.append(0.93)
    if s:
        if _contains(e, s):
            scores.append(0.90)
        scores.append(0.85 * SequenceMatcher(None, s, e).ratio())
        tokens = [t for t in s.split() if len(t) > 1]
        if tokens and all(_contains(e, t) for t in tokens):
            scores.append(0.88)
    if f:
        scores.append(0.95 * SequenceMatcher(None, f, e).ratio())
        # The ERP title is often just a prefix of the Excel name (the bracket
        # part dropped). If every ERP title word appears in the Excel name,
        # treat it as a strong match.
        f_tokens = {_stem(token) for token in f.split()}
        e_tokens = [token for token in e.split() if len(token) > 1]
        if len(e_tokens) >= 3 and all(_stem(token) in f_tokens for token in e_tokens):
            scores.append(0.90)
    return max(scores)


@dataclass
class Match:
    full_name: str
    short: str
    project: ErpProject | None
    score: float
    note: str = ""


def match_projects(
    excel_names: list[str],
    erp_projects: list[ErpProject],
    manual: dict[str, str] | None = None,
    threshold: float = 0.72,
) -> list[Match]:
    manual = manual or {}
    by_id = {p.pid: p for p in erp_projects}
    matches: list[Match] = []

    for full in excel_names:
        code = short_name(full)

        manual_key = None
        for key in (code, full):
            if key in manual:
                manual_key = key
                break
        if manual_key:
            pid = str(manual[manual_key])
            project = by_id.get(pid)
            if project is None:
                project = ErpProject(pid=pid, name=f"(manual id {pid})")
            matches.append(Match(full, code, project, 1.0, "manual"))
            continue

        # Secondary key breaks ties between look-alike names (Learning Site I
        # vs II) when the ERP title drops the short code.
        ranked = sorted(
            (
                (
                    name_similarity(code, full, p.name, p.short),
                    SequenceMatcher(None, norm(full), norm(p.name)).ratio(),
                    p,
                )
                for p in erp_projects
            ),
            key=lambda item: (item[0], item[1]),
            reverse=True,
        )
        if not ranked:
            matches.append(Match(full, code, None, 0.0, "no projects found"))
            continue

        best_score, best_ratio, best_project = ranked[0]
        note = ""
        if len(ranked) > 1:
            second_score, second_ratio, second_project = ranked[1]
            if best_score - second_score < 0.03 and abs(best_ratio - second_ratio) < 0.02:
                note = f"ambiguous (also {second_project.name!r})"
        if best_score < threshold:
            matches.append(Match(
                full, code, None, best_score,
                note or f"best candidate {best_project.name!r} scored {best_score:.2f}",
            ))
        else:
            matches.append(Match(full, code, best_project, best_score, note))
    return matches


# --------------------------------------------------------------------------- #
# Export discovery (Activities page)
# --------------------------------------------------------------------------- #

@dataclass
class ExportCandidate:
    url: str
    method: str = "GET"
    data: dict = field(default_factory=dict)
    label: str = ""


_URL_IN_SCRIPT_RE = re.compile(
    r"""["']((?:https?://[^"'\s<>]+|/[^"'\s<>]*?)export[^"'\s<>]*)["']""", re.I
)
_ANCHOR_RE = re.compile(r"<a\b([^>]*)>(.*?)</a>", re.S | re.I)
_FORM_RE = re.compile(r"<form\b([^>]*)>(.*?)</form>", re.S | re.I)
_BUTTON_RE = re.compile(r"<button\b([^>]*)>(.*?)</button>", re.S | re.I)
_ATTR_RE = re.compile(r'([\w:-]+)\s*=\s*"([^"]*)"')


def _attrs(raw: str) -> dict:
    return {key.lower(): unescape(value) for key, value in _ATTR_RE.findall(raw)}


def find_activities_url(html: str, page_url: str, project_id: str) -> str:
    """Locate the Activities tab (``/projects/<id>/show``) from the dashboard."""
    fallback = urllib.parse.urljoin(page_url, f"/projects/{project_id}/show")
    for href, text in re.findall(r'<a\b[^>]*href="([^"]+)"[^>]*>(.*?)</a>', html, re.S | re.I):
        url = urllib.parse.urljoin(page_url, unescape(href))
        if f"/projects/{project_id}" not in url:
            continue
        label = strip_tags(text).lower()
        if label == "activities" or url.rstrip("/").endswith("/show") or "/activities" in url:
            return url
    return fallback


def find_export_candidates(html: str, page_url: str, token: str = "") -> list[ExportCandidate]:
    candidates: list[ExportCandidate] = []

    def add(url: str, method: str = "GET", data: dict | None = None, label: str = "") -> None:
        if not url:
            return
        url = urllib.parse.urljoin(page_url, unescape(url).replace("\\/", "/"))
        if not url.lower().startswith("http"):
            return
        for existing in candidates:
            if existing.method == method and existing.url == url:
                return
        candidates.append(ExportCandidate(url=url, method=method, data=data or {}, label=label))

    # 1) Links / buttons carrying an export URL.
    for raw_attrs, text in _ANCHOR_RE.findall(html) + _BUTTON_RE.findall(html):
        attrs = _attrs(raw_attrs)
        label = strip_tags(text) or attrs.get("title", "")
        blob = " ".join([label, attrs.get("href", ""), attrs.get("onclick", "")] +
                        [v for k, v in attrs.items() if k.startswith("data-")])
        if "export" not in blob.lower():
            continue
        url = attrs.get("href") or ""
        if not url or url.lower().startswith(("#", "javascript:")):
            url = next(
                (attrs[k] for k in ("data-url", "data-href", "data-link", "data-export-url",
                                    "data-action", "data-target-url") if attrs.get(k)),
                "",
            )
        if not url:
            onclick = attrs.get("onclick", "")
            match = re.search(r"""["']([^"']*export[^"']*)["']""", onclick, re.I)
            url = match.group(1) if match else ""
        if url:
            add(url, "GET", None, label)

    # 2) Forms whose action mentions export.
    for raw_attrs, body in _FORM_RE.findall(html):
        attrs = _attrs(raw_attrs)
        action = attrs.get("action", "")
        if "export" not in action.lower():
            continue
        data = {}
        for input_match in re.finditer(r"<input\b([^>]*)>", body, re.I):
            fields = _attrs(input_match.group(1))
            name = fields.get("name")
            if name:
                data[name] = fields.get("value", "")
        if token and "_token" not in data:
            data["_token"] = token
        add(action or page_url, attrs.get("method", "GET").upper(), data, "form")

    # 3) URLs hard-coded in inline scripts.
    for url in _URL_IN_SCRIPT_RE.findall(html):
        if "export" in url.lower():
            add(url, "GET", None, "inline script")

    # Prefer the plain "Export" button over "Export Activity", then Excel hints.
    def rank(candidate: ExportCandidate) -> tuple:
        label = norm(candidate.label)
        if label == "export":
            return (0, 0)
        if "export" in label and "activit" not in label:
            return (1, 0)
        if re.search(r"xls|excel", candidate.url + candidate.label, re.I):
            return (2, 0)
        return (3, 0)

    candidates.sort(key=rank)
    return candidates


def _export_labels(html: str) -> list[str]:
    """Visible labels of anchors/buttons that mention 'export' (for diagnostics)."""
    labels: list[str] = []
    for _attrs_raw, text in _ANCHOR_RE.findall(html) + _BUTTON_RE.findall(html):
        label = strip_tags(text)
        if label and "export" in label.lower() and label not in labels:
            labels.append(label[:40])
    return labels[:5]


def looks_like_file(response: Response) -> bool:
    content_type = response.header("content-type").lower()
    disposition = response.header("content-disposition").lower()
    if "attachment" in disposition:
        return True
    if any(key in content_type for key in ("spreadsheet", "excel", "octet-stream", "csv", "zip")):
        return True
    head = response.content[:8]
    if head.startswith(b"PK\x03\x04") or head.startswith(b"\xd0\xcf\x11\xe0"):
        return True
    text_head = response.content[:500].lstrip().lower()
    if text_head.startswith(b"<"):
        return False
    return False


def sniff_extension(response: Response) -> str:
    disposition = response.header("content-disposition")
    match = re.search(r'filename\*?=(?:UTF-8\'\')?"?([^";]+)"?', disposition, re.I)
    if match:
        suffix = Path(clean_text(match.group(1))).suffix.lower()
        if suffix in {".xlsx", ".xlsm", ".xls", ".csv", ".zip"}:
            return suffix
    content_type = response.header("content-type").lower()
    if "csv" in content_type:
        return ".csv"
    if "spreadsheetml" in content_type:
        return ".xlsx"
    if "ms-excel" in content_type:
        return ".xls"
    if "zip" in content_type:
        return ".zip"

    head = response.content[:8]
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return ".xls"
    if head.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
                names = archive.namelist()
            if any(name.startswith("xl/") for name in names):
                return ".xlsx"
            if any(name.lower().endswith(".csv") for name in names):
                return ".zip"
        except zipfile.BadZipFile:
            pass
        return ".zip"
    if b"," in response.content[:2000] or b";" in response.content[:2000]:
        return ".csv"
    return ".dat"


def download_export(
    client: ErpClient,
    project: ErpProject,
    short: str,
    out_dir: Path,
    referer: str,
    token: str,
    reuse: bool,
    keep_html: bool = False,
) -> tuple[Path | None, str]:
    """Return (file path, status message)."""
    existing = sorted(out_dir.glob(f"{safe_filename(short)}_{project.pid}.*"))
    if reuse and existing:
        return existing[0], "reused existing file"

    activities_url = urllib.parse.urljoin(client.base + "/", f"/projects/{project.pid}/show")
    dashboard_url = urllib.parse.urljoin(client.base + "/", f"/projects/{project.pid}/dashboard")

    pages = []
    for url in (dashboard_url, activities_url):
        response = client.get(url, referer=referer)
        if _is_login_page(response):
            raise AuthError("Session expired while opening the activities page.")
        if response.status < 400:
            pages.append((url, response))

    if not pages:
        return None, "dashboard/activities page not reachable"

    activities_url = find_activities_url(pages[0][1].text, pages[0][1].url, project.pid)
    activities_page = pages[-1][1]
    if activities_page.url != activities_url:
        response = client.get(activities_url, referer=pages[-1][0])
        if response.status < 400:
            activities_page = response

    attempts: list[str] = []
    html_dir = out_dir / "html"
    if keep_html:
        html_dir.mkdir(parents=True, exist_ok=True)
        for kind, (page_url, page) in zip(("dashboard", "activities"), pages):
            (html_dir / f"{project.pid}_{kind}.html").write_text(page.text, encoding="utf-8")

    for page_url, page in [(activities_url, activities_page)] + pages:
        candidates = find_export_candidates(page.text, page.url, token)
        if not candidates:
            labels = _export_labels(page.text)
            attempts.append(
                f"{page.url}: no export URL found"
                + (f" (buttons: {', '.join(labels)})" if labels else " (no Export buttons at all)")
            )
        if keep_html and not candidates:
            (html_dir / f"{project.pid}_no_export_found.html").write_text(
                page.text, encoding="utf-8"
            )
        for candidate in candidates:
            try:
                response = client.request(
                    candidate.method,
                    candidate.url,
                    data=candidate.data or None,
                    referer=page_url,
                )
                # Some ERP export endpoints reject a bare GET (CSRF / method
                # not allowed) but accept the same URL as a POST with the token.
                if (candidate.method == "GET" and response.status in (403, 405, 419)
                        and token):
                    response = client.post(
                        candidate.url,
                        {"_token": token},
                        referer=page_url,
                    )
            except ErpError as exc:
                attempts.append(f"{candidate.url}: {exc}")
                continue
            if _is_login_page(response):
                raise AuthError("Session expired while downloading the export.")
            if looks_like_file(response):
                extension = sniff_extension(response)
                target = out_dir / f"{safe_filename(short)}_{project.pid}{extension}"
                target.write_bytes(response.content)
                return target, f"downloaded {target.name} ({len(response.content):,} bytes)"
            attempts.append(f"{candidate.method} {candidate.url}: HTTP {response.status}, not a file")

    detail = "; ".join(attempts[:3]) if attempts else "no Export control found on the page"
    return None, f"export not found ({detail})"


# --------------------------------------------------------------------------- #
# Reading downloaded files
# --------------------------------------------------------------------------- #

def _decode_bytes(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def _is_blank(value) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    return isinstance(value, str) and not value.strip()


def _dedupe_columns(columns) -> list[str]:
    seen: dict[str, int] = {}
    result = []
    for index, column in enumerate(columns, start=1):
        name = "" if _is_blank(column) else str(column).strip()
        name = re.sub(r"\s+", " ", name)
        if not name:
            name = f"Column {index}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        result.append(name)
    return result


def _frame_from_raw(raw: pd.DataFrame) -> pd.DataFrame:
    if raw.empty:
        return pd.DataFrame()

    ncols = max(1, raw.shape[1])
    header_row = None
    for index in range(min(15, len(raw))):
        values = [v for v in raw.iloc[index].tolist() if not _is_blank(v)]
        if len(values) < max(3, int(ncols * 0.5)):
            continue
        textual = sum(isinstance(v, str) for v in values)
        if textual >= len(values) * 0.5:
            header_row = index
            break
    if header_row is None:
        header_row = 0

    header = raw.iloc[header_row].tolist()
    frame = raw.iloc[header_row + 1 :].copy()
    frame.columns = _dedupe_columns(header)
    frame = frame.dropna(how="all")
    frame = frame.loc[:, ~frame.apply(lambda column: column.map(_is_blank).all())]
    return frame.reset_index(drop=True)


def read_table(path: Path) -> pd.DataFrame:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        raw = pd.read_excel(path, sheet_name=0, header=None, dtype=object)
        return _frame_from_raw(raw)
    if suffix == ".xls":
        try:
            raw = pd.read_excel(path, sheet_name=0, header=None, dtype=object)
        except ImportError as exc:  # legacy .xls needs xlrd
            raise ValueError(f"{path.name} is a legacy .xls file - run: pip install xlrd") from exc
        return _frame_from_raw(raw)
    if suffix == ".csv":
        text = _decode_bytes(path.read_bytes())
        return _read_csv(text)
    if suffix == ".zip":
        frames = []
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if name.lower().endswith(".csv"):
                    frames.append(_read_csv(_decode_bytes(archive.read(name))))
                elif name.lower().endswith((".xlsx", ".xlsm")):
                    frames.append(_frame_from_raw(
                        pd.read_excel(io.BytesIO(archive.read(name)), sheet_name=0,
                                      header=None, dtype=object)
                    ))
        if not frames:
            raise ValueError(f"no CSV/XLSX inside {path.name}")
        return pd.concat(frames, ignore_index=True, sort=False)
    raise ValueError(f"unsupported export format: {path.suffix}")


def _read_csv(text: str) -> pd.DataFrame:
    separator = None
    try:
        import csv as _csv

        separator = _csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
    except Exception:
        separator = None
    raw = pd.read_csv(
        io.StringIO(text),
        header=None,
        dtype=object,
        sep=separator or ",",
        engine="python",
        on_bad_lines="skip",
        skip_blank_lines=False,
    )
    return _frame_from_raw(raw)


# --------------------------------------------------------------------------- #
# Excel workbook handling
# --------------------------------------------------------------------------- #

def read_project_names(workbook: Path, sheet_name: str) -> tuple[list[str], int, int]:
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise SystemExit(
            f"Sheet {sheet_name!r} not found. Available sheets: {', '.join(wb.sheetnames)}"
        )
    ws = wb[sheet_name]

    header_row, column = None, None
    for row in ws.iter_rows(min_row=1, max_row=min(15, ws.max_row or 1)):
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.strip().lower() == "project":
                header_row, column = cell.row, cell.column
                break
        if column:
            break
    if column is None:
        for row in ws.iter_rows(min_row=1, max_row=min(15, ws.max_row or 1)):
            for cell in row:
                value = cell.value
                if isinstance(value, str) and "project" in value.lower() and len(value) < 60:
                    header_row, column = cell.row, cell.column
                    break
            if column:
                break
    if column is None:
        raise SystemExit(f"Could not find a 'Project' column in sheet {sheet_name!r}.")

    names: list[str] = []
    seen: set[str] = set()
    for row_index in range(header_row + 1, (ws.max_row or header_row) + 1):
        value = ws.cell(row=row_index, column=column).value
        if not isinstance(value, str):
            continue
        text = clean_text(value)
        key = text.lower()
        if not text or key in seen:
            continue
        seen.add(key)
        names.append(text)

    wb.close()
    return names, header_row, column


def write_sheet(ws, frame: pd.DataFrame) -> None:
    ws.append(list(frame.columns))
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN

    for row in frame.itertuples(index=False):
        ws.append([_python_value(value) for value in row])

    if frame.shape[0]:
        ws.auto_filter.ref = f"A1:{get_column_letter(frame.shape[1])}{frame.shape[0] + 1}"
    ws.freeze_panes = "A2"

    for index, column in enumerate(frame.columns, start=1):
        longest = len(str(column))
        for value in frame[column].head(400):
            if not _is_blank(value):
                longest = max(longest, len(str(value)))
        ws.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 10), 45)
        ws.column_dimensions[get_column_letter(index)].alignment = Alignment(vertical="top")


def _python_value(value):
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if hasattr(value, "to_pydatetime"):
        try:
            return value.to_pydatetime()
        except Exception:
            pass
    if hasattr(value, "item") and not isinstance(value, (str, bytes)):
        try:
            value = value.item()
        except Exception:
            pass
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, (list, dict, set, tuple)):
        return str(value)
    return value


def save_workbook(wb, path: Path) -> None:
    temporary = path.with_name(f"{path.stem}.tmp{path.suffix}")
    wb.save(temporary)
    os.replace(temporary, path)


# --------------------------------------------------------------------------- #
# Main flow
# --------------------------------------------------------------------------- #

def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Download all ERP project activity exports and merge them into Sheet3.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    parser.add_argument("--workbook", default=DEFAULT_WORKBOOK, help="Path to the .xlsx file")
    parser.add_argument("--sheet", default="Sheet2", help="Sheet holding the project list")
    parser.add_argument("--out-sheet", default="Sheet3", help="Sheet to write the merged data")
    parser.add_argument("--exports-dir", default="erp_exports", help="Folder for downloaded files")
    parser.add_argument("--base-url", default=BASE_URL, help="ERP base URL")
    parser.add_argument("--user", default=os.environ.get("HERD_ERP_USER"), help="ERP username")
    parser.add_argument("--password", default=os.environ.get("HERD_ERP_PASS"), help="ERP password")
    parser.add_argument("--map-file", default=None, help="JSON file with manual short-name -> id entries")
    parser.add_argument("--list-only", action="store_true", help="Only print the project matching")
    parser.add_argument("--reuse", action="store_true", help="Reuse files already in --exports-dir")
    parser.add_argument("--keep-html", action="store_true", help="Save page HTML for debugging")
    parser.add_argument("--no-backup", action="store_true", help="Do not create a backup copy")
    parser.add_argument("--no-map-sheet", action="store_true", help="Do not write the 'Project Map' sheet")
    parser.add_argument("--insecure", action="store_true", help="Skip TLS certificate verification")
    parser.add_argument("--delay", type=float, default=0.4, help="Pause between requests (seconds)")
    parser.add_argument("--limit", type=int, default=0, help="Only handle the first N projects")
    parser.add_argument("--strict", action="store_true", help="Exit non-zero when a project fails")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    log(f"erp_export.py build {VERSION}")
    workbook = Path(args.workbook).expanduser()

    if not workbook.exists():
        log(f"ERROR: workbook not found: {workbook}")
        return 1

    log(f"Workbook : {workbook}")
    names, header_row, column = read_project_names(workbook, args.sheet)
    if not names:
        log(f"ERROR: no project names found in {args.sheet!r} column {column}.")
        return 1
    log(f"Sheet    : {args.sheet!r}, 'Project' in column "
        f"{get_column_letter(column)}, header row {header_row}")
    log(f"Projects : {len(names)}\n")

    if args.limit:
        names = names[: args.limit]

    manual = {}
    if args.map_file:
        try:
            manual = json.loads(Path(args.map_file).read_text(encoding="utf-8-sig"))
        except (OSError, json.JSONDecodeError) as exc:
            log(f"ERROR: could not read --map-file {args.map_file}: {exc}")
            return 1
        log(f"Loaded {len(manual)} manual mapping(s) from {args.map_file}")

    username = args.user
    password = args.password
    if not username:
        username = input("ERP username: ").strip()
    if not password:
        import getpass

        password = getpass.getpass("ERP password: ")

    client = ErpClient(args.base_url, insecure=args.insecure, delay=args.delay)
    client.login(username, password)

    log("Reading /projects ...")
    erp_projects = scrape_projects(client)
    if not erp_projects:
        log("ERROR: no projects could be read from /projects. "
            "The page layout may differ from what this script expects.")
        return 1
    log(f"Found {len(erp_projects)} project(s) in the ERP.\n")

    matches = match_projects(names, erp_projects, manual)

    log(f"{'#':>3}  {'SHORT NAME':<34} {'SCORE':>5}  ERP PROJECT")
    log("-" * 110)
    for index, match in enumerate(matches, start=1):
        target = match.project.name if match.project else "-- NOT FOUND --"
        log(f"{index:>3}  {match.short[:34]:<34} {match.score:>5.2f}  {target[:60]}")
        if match.note:
            log(f"       note: {match.note}")
    log("")

    missing = [m for m in matches if not m.project]
    if missing:
        log(f"WARNING: {len(missing)} project(s) could not be matched. "
            "Use --map-file to map them manually.\n")

    if args.list_only:
        return 0

    exports_dir = Path(args.exports_dir).expanduser()
    exports_dir.mkdir(parents=True, exist_ok=True)

    token = _first(client.get("/projects").text,
                   r'name="csrf-token"\s+content="([^"]+)"',
                   r'name="_token"\s+value="([^"]+)"')

    downloaded: list[tuple[Match, Path]] = []
    status_rows: list[dict] = []

    for index, match in enumerate(matches, start=1):
        if match.project is None:
            status_rows.append(_status(match, "unmatched", "", "", ""))
            continue
        log(f"[{index}/{len(matches)}] {match.short} -> {match.project.name} (id {match.project.pid})")
        try:
            path, message = download_export(
                client, match.project, match.short, exports_dir,
                referer=urllib.parse.urljoin(client.base + "/", "/projects"),
                token=token, reuse=args.reuse, keep_html=args.keep_html,
            )
        except AuthError:
            raise
        except Exception as exc:  # noqa: BLE001 - keep going for the other projects
            path, message = None, f"error: {type(exc).__name__}: {exc}"

        log(f"      {message}")
        if path:
            downloaded.append((match, path))
            status_rows.append(_status(match, "ok", path.name, "", ""))
        else:
            status_rows.append(_status(match, "failed", "", "", message))

    if not downloaded:
        log("\nNothing was downloaded - no Sheet3 was written.")
        return 1

    log(f"\nReading {len(downloaded)} export file(s) ...")
    frames = []
    for match, path in downloaded:
        try:
            frame = read_table(path)
        except Exception as exc:  # noqa: BLE001
            log(f"  ! {path.name}: {type(exc).__name__}: {exc}")
            continue
        if frame.empty:
            log(f"  ! {path.name}: no data rows")
            continue
        log(f"  + {path.name}: {len(frame):,} rows x {frame.shape[1]} columns")
        frame.insert(0, "Project (Full)", match.project.name)
        frame.insert(0, "Project (Short)", match.short)
        frames.append(frame)
        for row in status_rows:
            if row["Short"] == match.short:
                row["Rows"] = len(frame)

    if not frames:
        log("No usable rows were read - no Sheet3 was written.")
        return 1

    merged = pd.concat(frames, ignore_index=True, sort=False)
    merged = merged.dropna(how="all")
    log(f"\nMerged: {len(merged):,} rows x {merged.shape[1]} columns")

    log("Writing workbook ...")
    if not args.no_backup:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = workbook.with_name(f"{workbook.stem}_backup_{stamp}{workbook.suffix}")
        shutil.copy2(workbook, backup)
        log(f"  backup: {backup.name}")

    wb = openpyxl.load_workbook(workbook, data_only=False, keep_links=True)
    if args.out_sheet in wb.sheetnames:
        del wb[args.out_sheet]
    ws = wb.create_sheet(args.out_sheet)
    write_sheet(ws, merged)

    if not args.no_map_sheet:
        map_name = "Project Map"
        if map_name in wb.sheetnames:
            del wb[map_name]
        map_ws = wb.create_sheet(map_name)
        map_frame = pd.DataFrame(status_rows, columns=[
            "Project (as in Excel)", "Short", "ERP Project", "ERP ID", "Status", "File", "Rows", "Note",
        ])
        write_sheet(map_ws, map_frame)

    saved = False
    for attempt in range(1, 4):
        try:
            save_workbook(wb, workbook)
            saved = True
            break
        except PermissionError:
            log(f"  workbook is locked (open in Excel?) - retry {attempt}/3 in 6s ...")
            time.sleep(6)
    if not saved:
        log(f"ERROR: could not save {workbook.name} - close it in Excel and run again.")
        log("       Your downloads are safe in the exports folder; use --reuse to skip them.")
        return 1

    # Confirm the sheet actually made it to disk, including the R4R-style rows.
    try:
        check = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
        written = max((check[args.out_sheet].max_row or 1) - 1, 0)
        check.close()
        log(f"  verified on disk: {args.out_sheet} has {written:,} data rows "
            f"(expected {len(merged):,})")
        if written != len(merged):
            log("  WARNING: row count differs from what was merged - check the file.")
    except Exception as exc:  # noqa: BLE001
        log(f"  (could not re-open the workbook to verify: {type(exc).__name__}: {exc})")

    log(f"  saved sheet {args.out_sheet!r} to {workbook}")

    ok = sum(1 for row in status_rows if row["Status"] == "ok")
    log(f"\nDone: {ok}/{len(matches)} projects exported.")
    if ok < len(matches):
        for row in status_rows:
            if row["Status"] != "ok":
                log(f"  - {row['Short']}: {row['Status']} {row['Note']}")
        if args.strict:
            return 2
    return 0


def _status(match: Match, status: str, filename: str, rows, note: str) -> dict:
    return {
        "Project (as in Excel)": match.full_name,
        "Short": match.short,
        "ERP Project": match.project.name if match.project else "",
        "ERP ID": match.project.pid if match.project else "",
        "Status": status,
        "File": filename,
        "Rows": rows,
        "Note": note,
    }


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    sys.exit(main())
